"""FastAPI service for ASAP2 parsing and experiment generation."""

from __future__ import annotations

import io
import json
import sys
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List

_ROOT = Path(__file__).resolve().parents[2]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from a2l_parser import build_experiments, parse_a2l_text
from docx import Document
from docx.shared import Pt
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from openpyxl import Workbook

app = FastAPI(title="A2L Engineering Workbench API", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/api/health")
def health() -> Dict[str, str]:
    return {"status": "ok", "service": "a2l-workbench"}


@app.post("/api/parse")
async def parse_upload(file: UploadFile = File(...)) -> JSONResponse:
    if not file.filename or not file.filename.lower().endswith((".a2l", ".aml", ".txt")):
        raise HTTPException(status_code=400, detail="Expected an .a2l / .aml / .txt ASAP2 file")
    raw = await file.read()
    try:
        text = raw.decode("utf-8", errors="replace")
    except Exception as exc:  # pragma: no cover
        raise HTTPException(status_code=400, detail=f"Could not decode file: {exc}") from exc
    try:
        data = parse_a2l_text(text)
        experiments = build_experiments(data)
        payload = {"dataset": data, "experiments": experiments, "sourceFile": file.filename}
        return JSONResponse(payload)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Parse error: {exc}") from exc


def _build_word_report(dataset: Dict[str, Any], experiments: List[Dict[str, Any]]) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    meta = dataset.get("meta", {})
    doc.add_heading("A2L Engineering Workbench — Technical Report", level=0)
    doc.add_paragraph(f"Institution: National Institute of Technology Karnataka, Surathkal")
    doc.add_paragraph(f"Project: {meta.get('projectName','')}")
    doc.add_paragraph(f"ECU / Module: {meta.get('ecuName','')}")
    doc.add_paragraph(f"ASAP2 version: {meta.get('asap2Version','')}")
    doc.add_paragraph(f"Generated (UTC): {datetime.utcnow().isoformat()}Z")

    doc.add_heading("Dataset summary", level=1)
    counts = dataset.get("counts", {})
    for k, v in counts.items():
        doc.add_paragraph(f"{k}: {v}", style="List Bullet")

    doc.add_heading("Selected experiments", level=1)
    for ex in experiments[:40]:
        doc.add_heading(f"Experiment {ex.get('id')}: {ex.get('title')}", level=2)
        doc.add_paragraph(f"Subsystem: {ex.get('subsystem')} | Difficulty: {ex.get('difficulty')}")
        doc.add_paragraph(ex.get("engineeringPurpose", ""))
        for step in ex.get("incaProcedure", [])[:25]:
            doc.add_paragraph(step, style="List Number")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def _build_excel_workbook(dataset: Dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Characteristics"
    headers = [
        "Name",
        "Description",
        "Address",
        "Type",
        "Lower",
        "Upper",
        "COMPU",
        "Memory",
        "Function",
    ]
    ws.append(headers)
    for c in dataset.get("characteristics", []):
        ws.append(
            [
                c.get("name"),
                c.get("description"),
                c.get("address"),
                c.get("type"),
                c.get("lowerLimit"),
                c.get("upperLimit"),
                c.get("compuMethod"),
                c.get("memorySegment"),
                c.get("relatedFunction"),
            ]
        )

    wm = wb.create_sheet("Measurements")
    wm.append(
        [
            "Name",
            "Description",
            "Datatype",
            "ECU Address",
            "COMPU",
            "Resolution",
            "Lower",
            "Upper",
        ]
    )
    for m in dataset.get("measurements", []):
        wm.append(
            [
                m.get("name"),
                m.get("description"),
                m.get("datatype"),
                m.get("ecuAddress"),
                m.get("compuMethod"),
                m.get("resolution"),
                m.get("lowerLimit"),
                m.get("upperLimit"),
            ]
        )

    wf = wb.create_sheet("Functions")
    wf.append(["Name", "Description", "Inputs", "Outputs", "Locals", "Calibrations"])
    for f in dataset.get("functions", []):
        wf.append(
            [
                f.get("name"),
                f.get("description"),
                ", ".join(f.get("inputs") or []),
                ", ".join(f.get("outputs") or []),
                ", ".join(f.get("localMeasurements") or []),
                ", ".join(f.get("calibrations") or []),
            ]
        )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


@app.post("/api/export/word")
async def export_word(payload: Dict[str, Any]) -> StreamingResponse:
    dataset = payload.get("dataset") or {}
    experiments = payload.get("experiments") or []
    data = _build_word_report(dataset, experiments)
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": 'attachment; filename="a2l_report.docx"'},
    )


@app.post("/api/export/excel")
async def export_excel(payload: Dict[str, Any]) -> StreamingResponse:
    dataset = payload.get("dataset") or {}
    data = _build_excel_workbook(dataset)
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="a2l_tables.xlsx"'},
    )


@app.post("/api/export/bundle")
async def export_bundle(payload: Dict[str, Any]) -> StreamingResponse:
    dataset = payload.get("dataset") or {}
    experiments = payload.get("experiments") or []
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("dataset.json", json.dumps(dataset, indent=2)[:12_000_000])
        zf.writestr("experiments.json", json.dumps(experiments, indent=2)[:12_000_000])
        zf.writestr("report.docx", _build_word_report(dataset, experiments))
        zf.writestr("tables.xlsx", _build_excel_workbook(dataset))
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/zip",
        headers={"Content-Disposition": 'attachment; filename="a2l_export_bundle.zip"'},
    )
