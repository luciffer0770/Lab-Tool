"""
Standalone offline ASAP2 / A2L inspection tool (Python + Tkinter).

National Institute of Technology Karnataka — Automotive Electronics Laboratory
"""

from __future__ import annotations

import json
import sys
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from a2l_parser import build_experiments, parse_a2l_text  # noqa: E402

try:
    from docx import Document
    from docx.shared import Pt

    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from openpyxl import Workbook

    HAS_XLSX = True
except ImportError:
    HAS_XLSX = False


class A2LWorkbenchTk(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("A2L Engineering Workbench — Standalone (NITK)")
        self.geometry("1200x800")
        self.configure(bg="#F5F7FA")
        self.dataset = None
        self.experiments = []
        self._build_ui()

    def _build_ui(self) -> None:
        top = tk.Frame(self, bg="#FFFFFF", padx=12, pady=10)
        top.pack(fill=tk.X)
        tk.Label(
            top,
            text="A2L Engineering Workbench",
            font=("Segoe UI", 14, "bold"),
            fg="#D71920",
            bg="#FFFFFF",
        ).pack(side=tk.LEFT)
        tk.Label(
            top,
            text="Automotive Electronics Laboratory — NITK Surathkal",
            font=("Segoe UI", 9),
            fg="#5A5A5A",
            bg="#FFFFFF",
        ).pack(side=tk.LEFT, padx=12)
        ttk.Button(top, text="Open A2L…", command=self._open_file).pack(side=tk.RIGHT)
        ttk.Button(top, text="Export JSON", command=self._export_json).pack(side=tk.RIGHT, padx=4)
        ttk.Button(top, text="Export Excel", command=self._export_xlsx).pack(side=tk.RIGHT, padx=4)
        ttk.Button(top, text="Export Word", command=self._export_docx).pack(side=tk.RIGHT, padx=4)

        self.nb = ttk.Notebook(self)
        self.nb.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)

        self.tab_summary = tk.Frame(self.nb, bg="#F5F7FA")
        self.tab_chars = tk.Frame(self.nb, bg="#F5F7FA")
        self.tab_meas = tk.Frame(self.nb, bg="#F5F7FA")
        self.tab_func = tk.Frame(self.nb, bg="#F5F7FA")
        self.tab_exp = tk.Frame(self.nb, bg="#F5F7FA")
        self.tab_raw = tk.Frame(self.nb, bg="#F5F7FA")

        for t, title in [
            (self.tab_summary, "Summary"),
            (self.tab_chars, "Characteristics"),
            (self.tab_meas, "Measurements"),
            (self.tab_func, "Functions"),
            (self.tab_exp, "Experiments"),
            (self.tab_raw, "Raw preview"),
        ]:
            self.nb.add(t, text=title)

        self.summary_text = tk.Text(self.tab_summary, wrap="word", font=("Segoe UI", 10), height=20)
        self.summary_text.pack(fill=tk.BOTH, expand=True, padx=6, pady=6)

        self._make_list_tab(self.tab_chars)
        self._make_list_tab(self.tab_meas)
        self._make_func_tab()
        self._make_exp_tab()
        self.raw_text = tk.Text(self.tab_raw, wrap="none", font=("JetBrains Mono", 9))
        ys = ttk.Scrollbar(self.tab_raw, command=self.raw_text.yview)
        self.raw_text.configure(yscrollcommand=ys.set)
        self.raw_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        ys.pack(side=tk.RIGHT, fill=tk.Y)

    def _make_list_tab(self, parent: tk.Frame) -> None:
        cols_frame = tk.Frame(parent, bg="#F5F7FA")
        cols_frame.pack(fill=tk.BOTH, expand=True)
        tree = ttk.Treeview(cols_frame, show="headings")
        tree.pack(fill=tk.BOTH, expand=True, side=tk.LEFT)
        sb = ttk.Scrollbar(cols_frame, orient="vertical", command=tree.yview)
        sb.pack(side=tk.RIGHT, fill=tk.Y)
        tree.configure(yscrollcommand=sb.set)
        if parent == self.tab_chars:
            self.tree_chars = tree
        else:
            self.tree_meas = tree

    def _make_func_tab(self) -> None:
        paned = ttk.Panedwindow(self.tab_func, orient=tk.HORIZONTAL)
        paned.pack(fill=tk.BOTH, expand=True, padx=4, pady=4)
        left = ttk.Frame(paned)
        right = ttk.Frame(paned)
        paned.add(left, weight=1)
        paned.add(right, weight=2)
        self.func_list = tk.Listbox(left, exportselection=False, font=("Segoe UI", 10))
        self.func_list.pack(fill=tk.BOTH, expand=True)
        self.func_list.bind("<<ListboxSelect>>", self._on_func_select)
        self.func_detail = tk.Text(right, wrap="word", font=("Segoe UI", 10))
        self.func_detail.pack(fill=tk.BOTH, expand=True)

    def _make_exp_tab(self) -> None:
        paned = ttk.Panedwindow(self.tab_exp, orient=tk.HORIZONTAL)
        paned.pack(fill=tk.BOTH, expand=True, padx=4, pady=4)
        left = ttk.Frame(paned)
        right = ttk.Frame(paned)
        paned.add(left, weight=1)
        paned.add(right, weight=2)
        self.exp_list = tk.Listbox(left, exportselection=False, font=("Segoe UI", 9))
        self.exp_list.pack(fill=tk.BOTH, expand=True)
        self.exp_list.bind("<<ListboxSelect>>", self._on_exp_select)
        self.exp_detail = tk.Text(right, wrap="word", font=("Segoe UI", 10))
        self.exp_detail.pack(fill=tk.BOTH, expand=True)

    def _open_file(self) -> None:
        path = filedialog.askopenfilename(filetypes=[("A2L", "*.a2l"), ("Text", "*.txt"), ("All", "*.*")])
        if not path:
            return
        text = Path(path).read_text(encoding="utf-8", errors="replace")
        self.dataset = parse_a2l_text(text)
        self.experiments = build_experiments(self.dataset)
        self._refresh_all(Path(path).name)

    def _refresh_all(self, filename: str) -> None:
        d = self.dataset
        assert d is not None
        meta = d["meta"]
        cts = d["counts"]
        self.summary_text.delete("1.0", tk.END)
        self.summary_text.insert(
            tk.END,
            f"File: {filename}\n"
            f"ECU: {meta.get('ecuName')}\n"
            f"Project: {meta.get('projectName')}\n"
            f"ASAP2 version: {meta.get('asap2Version')}\n\n"
            f"Counts:\n{json.dumps(cts, indent=2)}\n\n"
            f"Experiments generated: {len(self.experiments)}\n",
        )

        for tree, rows, cols in [
            (self.tree_chars, d["characteristics"], ("name", "description", "address", "type", "compuMethod")),
            (self.tree_meas, d["measurements"], ("name", "description", "datatype", "ecuAddress", "compuMethod")),
        ]:
            tree.delete(*tree.get_children())
            tree["columns"] = cols
            for c in cols:
                tree.heading(c, text=c)
                tree.column(c, width=140 if c != "description" else 280)
            for r in rows:
                tree.insert("", tk.END, values=tuple(r.get(c, "") for c in cols))

        self.func_list.delete(0, tk.END)
        for f in d["functions"]:
            self.func_list.insert(tk.END, f["name"])
        self.func_detail.delete("1.0", tk.END)

        self.exp_list.delete(0, tk.END)
        for ex in self.experiments:
            self.exp_list.insert(tk.END, f"{ex['id']:04d} — {ex['title']}")
        self.exp_detail.delete("1.0", tk.END)

        self.raw_text.delete("1.0", tk.END)
        self.raw_text.insert(tk.END, d.get("rawPreview", "")[:200000])

    def _on_func_select(self, _evt=None) -> None:
        sel = self.func_list.curselection()
        if not sel or not self.dataset:
            return
        f = self.dataset["functions"][sel[0]]
        lines = [
            f"Name: {f['name']}",
            f"Description: {f.get('description','')}",
            "",
            f"Inputs ({len(f.get('inputs') or [])}):",
            "\n".join(f"  - {x}" for x in (f.get("inputs") or [])),
            "",
            f"Outputs ({len(f.get('outputs') or [])}):",
            "\n".join(f"  - {x}" for x in (f.get("outputs") or [])),
            "",
            f"Local measurements ({len(f.get('localMeasurements') or [])}):",
            "\n".join(f"  - {x}" for x in (f.get("localMeasurements") or [])),
            "",
            f"Calibrations ({len(f.get('calibrations') or [])}):",
            "\n".join(f"  - {x}" for x in (f.get("calibrations") or [])),
        ]
        self.func_detail.delete("1.0", tk.END)
        self.func_detail.insert(tk.END, "\n".join(lines))

    def _on_exp_select(self, _evt=None) -> None:
        sel = self.exp_list.curselection()
        if not sel or not self.experiments:
            return
        ex = self.experiments[sel[0]]
        self.exp_detail.delete("1.0", tk.END)
        self.exp_detail.insert(tk.END, json.dumps(ex, indent=2))

    def _export_json(self) -> None:
        if not self.dataset:
            messagebox.showinfo("Export", "Open a file first.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".json", filetypes=[("JSON", "*.json")])
        if not path:
            return
        Path(path).write_text(json.dumps({"dataset": self.dataset, "experiments": self.experiments}, indent=2), encoding="utf-8")
        messagebox.showinfo("Export", "Saved.")

    def _export_xlsx(self) -> None:
        if not self.dataset:
            messagebox.showinfo("Export", "Open a file first.")
            return
        if not HAS_XLSX:
            messagebox.showerror("Export", "Install openpyxl: pip install openpyxl")
            return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
        if not path:
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "Characteristics"
        ch = self.dataset["characteristics"]
        if ch:
            ws.append(list(ch[0].keys()))
            for r in ch:
                ws.append(list(r.values()))
        wm = wb.create_sheet("Measurements")
        me = self.dataset["measurements"]
        if me:
            wm.append(list(me[0].keys()))
            for r in me:
                wm.append(list(r.values()))
        wb.save(path)
        messagebox.showinfo("Export", "Excel saved.")

    def _export_docx(self) -> None:
        if not self.dataset:
            messagebox.showinfo("Export", "Open a file first.")
            return
        if not HAS_DOCX:
            messagebox.showerror("Export", "Install python-docx: pip install python-docx")
            return
        path = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word", "*.docx")])
        if not path:
            return
        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        doc.add_heading("A2L Engineering Workbench — Technical Report", 0)
        meta = self.dataset["meta"]
        doc.add_paragraph(f"ECU: {meta.get('ecuName')}")
        doc.add_paragraph(f"Project: {meta.get('projectName')}")
        doc.add_paragraph(f"Counts: {json.dumps(self.dataset['counts'])}")
        for ex in self.experiments[:30]:
            doc.add_heading(f"Experiment {ex['id']}: {ex['title']}", level=2)
            doc.add_paragraph(ex.get("engineeringPurpose", ""))
        doc.save(path)
        messagebox.showinfo("Export", "Word report saved.")


def main() -> None:
    app = A2LWorkbenchTk()
    app.mainloop()


if __name__ == "__main__":
    main()
